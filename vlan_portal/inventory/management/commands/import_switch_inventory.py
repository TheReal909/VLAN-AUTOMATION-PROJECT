import ipaddress
from collections import Counter, defaultdict
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from inventory.models import Facility, Switch


REQUIRED_COLUMNS = {
    "facility_code",
    "facility_name",
    "facility_status",
    "address",
    "city",
    "state",
    "zip",
    "switch_name",
    "hostname",
    "management_ip",
    "model_family",
    "fastiron_version",
    "closet_role",
}


class Command(BaseCommand):
    help = "Import selected facility and logical-switch inventory from an XLSX workbook."

    def add_arguments(self, parser):
        parser.add_argument("workbook", type=Path)
        parser.add_argument(
            "--facility",
            action="append",
            dest="facilities",
            required=True,
            help="Facility code to import; repeat this option to select multiple facilities.",
        )
        parser.add_argument(
            "--apply",
            action="store_true",
            help="Write validated records to the database. Without this flag, perform a dry run.",
        )

    def handle(self, *args, **options):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise CommandError("Install openpyxl before importing XLSX files.") from exc

        workbook_path = options["workbook"]
        if not workbook_path.is_file():
            raise CommandError(f"Workbook not found: {workbook_path}")
        selected_codes = {code.strip().upper() for code in options["facilities"]}
        if not selected_codes:
            raise CommandError("Select at least one facility code.")

        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        if "combined inventory" not in workbook.sheetnames:
            raise CommandError("Workbook must contain a 'combined inventory' sheet.")

        rows = self._read_rows(workbook["combined inventory"], selected_codes)
        grouped = self._validate(rows, selected_codes)
        if options["apply"]:
            self._import(grouped)
            mode = "Imported"
        else:
            mode = "Dry run: would import"

        self.stdout.write(f"{mode} {len(grouped)} facilities and {sum(map(len, grouped.values()))} logical switches.")
        for code, records in sorted(grouped.items()):
            self.stdout.write(f"  {code}: {len(records)} switches")
        if not options["apply"]:
            self.stdout.write(self.style.WARNING("No database changes made. Repeat with --apply to import."))
        else:
            self.stdout.write(self.style.SUCCESS("Selected switch inventory imported."))

    @staticmethod
    def _read_rows(sheet, selected_codes):
        iterator = sheet.iter_rows(values_only=True)
        header = next(iterator, None)
        if not header:
            raise CommandError("'combined inventory' is empty.")
        headers = [str(value).strip() if value is not None else "" for value in header]
        missing = sorted(REQUIRED_COLUMNS - set(headers))
        if missing:
            raise CommandError(f"Missing required columns: {', '.join(missing)}")
        rows = []
        for row_number, values in enumerate(iterator, 2):
            record = {name: values[index] if index < len(values) else None for index, name in enumerate(headers)}
            name = str(record.get("switch_name") or "").strip()
            address = str(record.get("management_ip") or "").strip().lower()
            if name.lower() in {"caption", "switch_name"} or address in {"ip address", "management_ip"}:
                continue
            code = str(record.get("facility_code") or "").strip().upper()
            if code not in selected_codes:
                continue
            record["facility_code"] = code
            record["switch_name"] = name
            record["hostname"] = str(record.get("hostname") or "").strip().lower()
            record["management_ip"] = address
            record["closet_role"] = str(record.get("closet_role") or "").strip().lower()
            record["model_family"] = str(record.get("model_family") or "").strip().upper()
            record["fastiron_version"] = str(record.get("fastiron_version") or "").strip()
            record["row_number"] = row_number
            rows.append(record)
        return rows

    @staticmethod
    def _validate(rows, selected_codes):
        errors = []
        grouped = defaultdict(list)
        for row in rows:
            grouped[row["facility_code"]].append(row)

        for code in sorted(selected_codes):
            if code not in grouped:
                errors.append(f"Facility {code}: no rows found in workbook.")

        valid_families = {choice[0] for choice in Switch.ModelFamily.choices}
        valid_roles = {choice[0] for choice in Switch.ClosetRole.choices}
        for code, records in grouped.items():
            if not code.startswith("F") or len(code) != 6 or not code[1:].isdigit():
                errors.append(f"Facility {code}: code must be F followed by five digits.")
            facility_names = {str(row.get("facility_name") or "").strip() for row in records}
            if len(facility_names) != 1 or not next(iter(facility_names)):
                errors.append(f"Facility {code}: missing or inconsistent facility names.")
            roles = Counter(row["closet_role"] for row in records)
            if roles[Switch.ClosetRole.MDF] != 1:
                errors.append(f"Facility {code}: expected exactly one MDF, found {roles[Switch.ClosetRole.MDF]}.")
            if roles[Switch.ClosetRole.IDF] < 2:
                errors.append(f"Facility {code}: expected at least two IDFs for this traversal test.")
            for row in records:
                label = f"row {row['row_number']} ({code}/{row['switch_name']})"
                if not row["switch_name"] or not row["hostname"]:
                    errors.append(f"{label}: switch name and hostname are required.")
                try:
                    ipaddress.ip_address(row["management_ip"])
                except ValueError:
                    errors.append(f"{label}: invalid management IP.")
                if row["closet_role"] not in valid_roles:
                    errors.append(f"{label}: closet role must be MDF or IDF.")
                if row["model_family"] not in valid_families:
                    errors.append(f"{label}: unsupported model family '{row['model_family']}'.")
            for field in ("switch_name", "hostname", "management_ip"):
                values = [row[field].lower() if field != "management_ip" else row[field] for row in records]
                if len(set(values)) != len(values):
                    errors.append(f"Facility {code}: duplicate {field} values.")
            if Facility.objects.filter(code=code).exists():
                errors.append(f"Facility {code} already exists in the database; refusing to overwrite it.")

        all_hostnames = [row["hostname"] for records in grouped.values() for row in records]
        all_ips = [row["management_ip"] for records in grouped.values() for row in records]
        if len(set(all_hostnames)) != len(all_hostnames):
            errors.append("Selected facilities contain duplicate hostnames.")
        if len(set(all_ips)) != len(all_ips):
            errors.append("Selected facilities contain duplicate management IPs.")
        if set(all_hostnames) & set(Switch.objects.values_list("hostname", flat=True)):
            errors.append("One or more hostnames already exist in the database.")
        if set(all_ips) & set(Switch.objects.values_list("management_ip", flat=True)):
            errors.append("One or more management IPs already exist in the database.")
        if errors:
            raise CommandError("Import blocked:\n- " + "\n- ".join(errors))
        return grouped

    @transaction.atomic
    def _import(self, grouped):
        for code, records in grouped.items():
            sample = records[0]
            address = ", ".join(
                part for part in (
                    str(sample.get("address") or "").strip(),
                    str(sample.get("city") or "").strip(),
                    str(sample.get("state") or "").strip(),
                    str(sample.get("zip") or "").strip(),
                ) if part
            )
            facility = Facility.objects.create(
                code=code,
                name=str(sample["facility_name"]).strip(),
                address=address,
                is_active=str(sample.get("facility_status") or "").strip().lower() == "active",
            )
            by_name = {row["switch_name"]: row for row in records}
            created_switches = {}
            for row in records:
                upstream_name = str(row.get("upstream_switch") or "").strip()
                upstream = created_switches.get(upstream_name)
                created_switches[row["switch_name"]] = Switch.objects.create(
                    facility=facility,
                    name=row["switch_name"],
                    hostname=row["hostname"],
                    management_ip=row["management_ip"],
                    closet_role=row["closet_role"],
                    upstream_switch=upstream,
                    model_family=row["model_family"],
                    fastiron_version=row["fastiron_version"],
                    location_note=str(row.get("solarwinds_location") or "").strip(),
                    is_active=facility.is_active,
                )
            unresolved = [name for name, row in by_name.items() if row.get("upstream_switch") and row["upstream_switch"] not in by_name]
            if unresolved:
                raise CommandError(f"Facility {code}: upstream references not imported: {', '.join(unresolved)}")
